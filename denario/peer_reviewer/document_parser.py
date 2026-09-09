import os
import re
import subprocess
from typing import Dict, Any, Optional

try:
    import docx
except ImportError:
    docx = None

class PaperParser:
    """
    Extrator e estruturador de texto para artigos científicos (.pdf, .docx, .tex, .md, .txt)
    """

    def __init__(self, filepath: str):
        self.filepath = os.path.abspath(filepath)
        if not os.path.exists(self.filepath):
            raise FileNotFoundError(f"Arquivo não encontrado: {self.filepath}")

    def parse(self) -> Dict[str, Any]:
        ext = os.path.splitext(self.filepath)[1].lower()
        if ext == ".docx":
            raw_text = self._parse_docx()
        elif ext == ".pdf":
            raw_text = self._parse_pdf()
        elif ext in [".tex", ".latex"]:
            raw_text = self._parse_tex()
        elif ext in [".txt", ".md"]:
            raw_text = self._parse_txt()
        else:
            raise ValueError(f"Formato não suportado: {ext}. Suportados: .docx, .pdf, .tex, .md, .txt")

        sections = self._extract_sections(raw_text)
        return {
            "filepath": self.filepath,
            "filename": os.path.basename(self.filepath),
            "format": ext,
            "raw_text": raw_text,
            "char_count": len(raw_text),
            "word_count": len(raw_text.split()),
            "sections": sections
        }

    def _parse_docx(self) -> str:
        if docx is None:
            raise ImportError("Biblioteca 'python-docx' não instalada. Execute: pip install python-docx")
        doc = docx.Document(self.filepath)
        full_text = []
        for p in doc.paragraphs:
            if p.text.strip():
                full_text.append(p.text.strip())
        for tbl in doc.tables:
            for row in tbl.rows:
                row_vals = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_vals:
                    full_text.append(" | ".join(row_vals))
        return "\n\n".join(full_text)

    def _parse_pdf(self) -> str:
        # Tenta utilizar pdftotext do sistema
        try:
            res = subprocess.run(["pdftotext", self.filepath, "-"], capture_output=True, text=True, check=True)
            if res.stdout.strip():
                return res.stdout.strip()
        except Exception:
            pass

        # Fallback para pypdf se instalado
        try:
            import pypdf
            reader = pypdf.PdfReader(self.filepath)
            pages = [page.extract_text() for page in reader.pages if page.extract_text()]
            return "\n\n".join(pages)
        except Exception as e:
            raise RuntimeError(f"Não foi possível extrair texto do PDF. Instale 'pdftotext' ou 'pypdf': {e}")

    def _parse_tex(self) -> str:
        with open(self.filepath, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read()
        # Remove comentários do LaTeX
        content = re.sub(r'(?<!\\)%.*', '', content)
        return content.strip()

    def _parse_txt(self) -> str:
        with open(self.filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read().strip()

    def _extract_sections(self, text: str) -> Dict[str, str]:
        """
        Extrai seções principais do artigo utilizando expressões regulares flexíveis.
        """
        sections = {}
        
        # Padrões comuns de seções em artigos científicos
        sec_patterns = [
            ("titulo", r"(?:^|\n)([^\n]+)\n+(?:RESUMO|ABSTRACT|INTRODUÇÃO|INTRODUCTION)"),
            ("resumo", r"(?:RESUMO|ABSTRACT)\s*[\:–\-]?\s*(.*?)(?=\n\s*(?:PALAVRAS-CHAVE|KEYWORDS|1\.|INTRODUÇÃO|INTRODUCTION))"),
            ("introducao", r"(?:1\.\s*INTRODUÇÃO|1\.\s*INTRODUCTION|INTRODUÇÃO|INTRODUCTION)\s*\n+(.*?)(?=\n\s*(?:2\.|FUNDAMENTAÇÃO|METODOLOGIA|METHODOLOGY))"),
            ("metodologia", r"(?:METODOLOGIA|METHODOLOGY|MATERIAIS E MÉTODOS|MATERIALS AND METHODS)\s*\n+(.*?)(?=\n\s*(?:RESULTADOS|RESULTS|4\.|DISCUSSÃO))"),
            ("resultados", r"(?:RESULTADOS|RESULTS|RESULTADOS E DISCUSSÃO|RESULTS AND DISCUSSION)\s*\n+(.*?)(?=\n\s*(?:5\.|CONCLUSÃO|CONCLUSION|REFERÊNCIAS|REFERENCES))"),
            ("conclusao", r"(?:CONCLUSÃO|CONCLUSION|CONCLUSÕES|CONCLUSIONS)\s*\n+(.*?)(?=\n\s*(?:AGRADECIMENTOS|ACKNOWLEDGMENTS|REFERÊNCIAS|REFERENCES))"),
            ("referencias", r"(?:REFERÊNCIAS|REFERENCES|REFERÊNCIAS BIBLIOGRÁFICAS)\s*\n+(.*)")
        ]
        
        for key, pat in sec_patterns:
            m = re.search(pat, text, re.IGNORECASE | re.DOTALL)
            if m:
                sections[key] = m.group(1).strip()
                
        return sections
